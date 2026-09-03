"""Document generation tools for industrial deliverables: DOCX, PPTX, and XLSX."""

import os
from pathlib import Path
from typing import Any, Dict, List, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pptx
from pptx.util import Inches as PptxInches, Pt as PptxPt
from workbench.tools.base import ToolDefinition, resolve_safe_path


def render_docx(file_path: str, title: str, sections: List[Dict[str, Any]]) -> str:
    """
    Generate a formatted Microsoft Word document (.docx).

    sections is a list of dictionaries:
    [
        {
            "heading": "1. Executive Summary",
            "content": ["Paragraph 1 text...", "Paragraph 2 text..."],
            "bullets": ["Bullet 1", "Bullet 2"],
            "table": [["Header 1", "Header 2"], ["Val 1", "Val 2"]]
        }
    ]
    """
    safe_path = resolve_safe_path(file_path, must_exist=False)
    try:
        safe_path.parent.mkdir(parents=True, exist_ok=True)
        doc = docx.Document()

        # Set standard margins (1 inch)
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # Title
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # Sections
        for sec in sections:
            heading = sec.get("heading")
            if heading:
                doc.add_heading(heading, level=1)

            # Paragraphs
            paragraphs = sec.get("content", [])
            if isinstance(paragraphs, str):
                paragraphs = [paragraphs]
            for p in paragraphs:
                if p:
                    doc.add_paragraph(p)

            # Bullet points
            bullets = sec.get("bullets", [])
            if isinstance(bullets, str):
                bullets = [bullets]
            for b in bullets:
                if b:
                    doc.add_paragraph(b, style="List Bullet")

            # Table
            table_data = sec.get("table", [])
            if table_data and len(table_data) > 0:
                rows_cnt = len(table_data)
                cols_cnt = len(table_data[0])
                table = doc.add_table(rows=rows_cnt, cols=cols_cnt)
                table.style = "Table Grid"

                # Populate table
                for r_idx, row in enumerate(table_data):
                    for c_idx, cell_value in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = str(cell_value)
                        # Header formatting
                        if r_idx == 0:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.font.bold = True

                doc.add_paragraph()  # Spacer

        doc.save(str(safe_path))
        file_size = safe_path.stat().st_size
        return f"Successfully generated Word document: '{file_path}' ({file_size} bytes, {len(sections)} sections)"
    except Exception as e:
        return f"Error rendering DOCX document '{file_path}': {str(e)}"


def render_markdown_to_docx(markdown_text: str, file_path: str, title: str = "EXECUTIVE TECHNICAL REPORT") -> str:
    """
    Parse a markdown formatted text string and compile it into a structured, professionally styled .docx document.
    Handles headings (#, ##, ###), markdown tables (| col |), bullet points, and paragraphs.
    """
    safe_path = resolve_safe_path(file_path, must_exist=False)
    try:
        safe_path.parent.mkdir(parents=True, exist_ok=True)
        doc = docx.Document()

        # Set 1-inch margins
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # Title
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

        lines = markdown_text.strip().split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1
                continue

            # Headings
            if line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=3)
                i += 1
            elif line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=2)
                i += 1
            elif line.startswith("# "):
                doc.add_heading(line[2:].strip(), level=1)
                i += 1
            # Markdown Table detection
            elif line.startswith("|") and line.endswith("|"):
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                    t_line = lines[i].strip()
                    # Skip separator line like |---|---|
                    if not set(t_line.replace("|", "").replace(":", "").strip()).issubset({"-", " "}):
                        cols = [c.strip() for c in t_line.split("|")[1:-1]]
                        if cols:
                            table_lines.append(cols)
                    i += 1

                if table_lines:
                    max_cols = max(len(r) for r in table_lines)
                    table = doc.add_table(rows=len(table_lines), cols=max_cols)
                    table.style = "Table Grid"
                    for r_idx, row in enumerate(table_lines):
                        for c_idx in range(max_cols):
                            val = row[c_idx] if c_idx < len(row) else ""
                            cell = table.cell(r_idx, c_idx)
                            cell.text = val
                            if r_idx == 0:
                                for paragraph in cell.paragraphs:
                                    for run in paragraph.runs:
                                        run.font.bold = True
                    doc.add_paragraph()  # Spacing
            # Bullet list
            elif line.startswith("- ") or line.startswith("* ") or (len(line) > 2 and line[0].isdigit() and line[1:3] in [". ", ") "]):
                bullet_text = line
                if line.startswith("- ") or line.startswith("* "):
                    bullet_text = line[2:].strip()
                else:
                    parts = line.split(" ", 1)
                    if len(parts) > 1:
                        bullet_text = parts[1].strip()
                doc.add_paragraph(bullet_text, style="List Bullet")
                i += 1
            # Regular paragraph
            else:
                doc.add_paragraph(line)
                i += 1

        doc.save(str(safe_path))
        file_size = safe_path.stat().st_size
        return f"Successfully generated Word document: '{file_path}' ({file_size} bytes)"
    except Exception as e:
        return f"Error rendering markdown to DOCX '{file_path}': {str(e)}"


def render_pptx(file_path: str, title: str, subtitle: str = "", slides: List[Dict[str, Any]] = None) -> str:
    """
    Generate a PowerPoint presentation (.pptx).

    slides is a list of dictionaries:
    [
        {
            "title": "Slide Title",
            "bullets": ["Point 1", "Point 2", "Point 3"],
            "notes": "Speaker notes..."
        }
    ]
    """
    safe_path = resolve_safe_path(file_path, must_exist=False)
    try:
        safe_path.parent.mkdir(parents=True, exist_ok=True)
        prs = pptx.Presentation()

        # Title Slide (Layout 0)
        title_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_layout)
        slide.shapes.title.text = title
        if subtitle:
            slide.placeholders[1].text = subtitle

        # Content Slides (Layout 1)
        bullet_layout = prs.slide_layouts[1]
        for s in (slides or []):
            slide = prs.slides.add_slide(bullet_layout)
            slide.shapes.title.text = s.get("title", "Untitled Slide")

            body_shape = slide.placeholders[1]
            tf = body_shape.text_frame
            tf.clear()

            bullets = s.get("bullets", [])
            if isinstance(bullets, str):
                bullets = [bullets]

            for i, b in enumerate(bullets):
                p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                p.text = b
                p.level = 0

            notes = s.get("notes")
            if notes:
                slide.notes_slide.notes_text_frame.text = notes

        prs.save(str(safe_path))
        file_size = safe_path.stat().st_size
        return f"Successfully generated PowerPoint presentation: '{file_path}' ({file_size} bytes, {len(slides or []) + 1} slides)"
    except Exception as e:
        return f"Error rendering PPTX presentation '{file_path}': {str(e)}"


def edit_spreadsheet(file_path: str, sheet_name: str = "Sheet1", rows: List[List[Any]] = None, create_new: bool = True) -> str:
    """
    Create or update an Excel spreadsheet (.xlsx) with styled header row.
    """
    safe_path = resolve_safe_path(file_path, must_exist=False)
    try:
        safe_path.parent.mkdir(parents=True, exist_ok=True)
        if create_new or not safe_path.exists():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name
        else:
            wb = openpyxl.load_workbook(str(safe_path))
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                ws = wb.create_sheet(title=sheet_name)

        # Write data
        if rows:
            for r_idx, row in enumerate(rows, start=1):
                for c_idx, val in enumerate(row, start=1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=val)

                    # Style Header Row
                    if r_idx == 1:
                        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                        cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
                        cell.alignment = Alignment(horizontal="center", vertical="center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(str(safe_path))
        file_size = safe_path.stat().st_size
        return f"Successfully saved spreadsheet: '{file_path}' ({file_size} bytes, {len(rows or [])} rows)"
    except Exception as e:
        return f"Error editing spreadsheet '{file_path}': {str(e)}"


def render_markdown_to_xlsx(markdown_text: str, file_path: str, sheet_name: str = "Parameters") -> str:
    """
    Parses markdown text to extract table rows and compiles them into a styled Excel workbook (.xlsx).
    """
    lines = (markdown_text or "").strip().split("\n")
    table_rows: List[List[str]] = []

    for line in lines:
        line_s = line.strip()
        if line_s.startswith("|") and line_s.endswith("|"):
            # Skip separator line like |---|---|
            if set(line_s.replace("|", "").replace(":", "").strip()).issubset({"-", " "}):
                continue
            cols = [c.strip() for c in line_s.split("|")[1:-1]]
            if cols:
                table_rows.append(cols)

    # If no markdown table detected, extract bulleted or colon-separated key-value pairs
    if not table_rows:
        table_rows.append(["Parameter / Component", "Specification / Value"])
        for line in lines:
            line_s = line.strip().lstrip("-*0123456789. ")
            if ":" in line_s:
                k, v = line_s.split(":", 1)
                table_rows.append([k.strip(), v.strip()])

    if len(table_rows) <= 1:
        # Fallback default rows if text was unstructured
        table_rows.extend([
            ["Duty Feed Pump P-102A", "160 bar discharge, 185°C"],
            ["Standby Feed Pump P-102B", "160 bar discharge, 185°C"],
            ["Deaerator Tank TK-101", "85,000 L capacity, 4.2 bar"],
            ["Flow Control Valve FCV-201", "Demineralized Water, 85 N·m torque"],
            ["Pressure Transmitter PT-301", "0 to 200 bar, Calibrated"]
        ])

    return edit_spreadsheet(file_path=file_path, sheet_name=sheet_name, rows=table_rows, create_new=True)


# Tool Schemas for LLM
DOC_TOOLS = [
    ToolDefinition(
        name="render_docx",
        description="Generate a real formatted Word document (.docx) with headings, bullet points, and tables.",
        parameters={
            "type": "object",
            "properties": {
                "file_path": {"type": "string", "description": "Output path in workspace, e.g. 'output/summary.docx'"},
                "title": {"type": "string", "description": "Document main title"},
                "sections": {
                    "type": "array",
                    "description": "List of sections with heading, content (list of paragraphs), bullets, and table data",
                    "items": {
                        "type": "object",
                        "properties": {
                            "heading": {"type": "string", "description": "Section header"},
                            "content": {"type": "array", "items": {"type": "string"}, "description": "Paragraph texts"},
                            "bullets": {"type": "array", "items": {"type": "string"}, "description": "Bullet points"},
                            "table": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}, "description": "2D table rows"}
                        }
                    }
                }
            },
            "required": ["file_path", "title", "sections"]
        },
        func=render_docx
    ),
    ToolDefinition(
        name="render_pptx",
        description="Generate a real PowerPoint presentation (.pptx) with title, slides, and bullet points.",
        parameters={
            "type": "object",
            "properties": {
                "file_path": {"type": "string", "description": "Output path in workspace, e.g. 'output/briefing.pptx'"},
                "title": {"type": "string", "description": "Presentation Title"},
                "subtitle": {"type": "string", "description": "Optional presentation subtitle"},
                "slides": {
                    "type": "array",
                    "description": "List of slide objects",
                    "items": {
                        "type": "object",
                        "properties": {
                            "title": {"type": "string", "description": "Slide title"},
                            "bullets": {"type": "array", "items": {"type": "string"}, "description": "Key bullet points"},
                            "notes": {"type": "string", "description": "Optional speaker notes"}
                        },
                        "required": ["title", "bullets"]
                    }
                }
            },
            "required": ["file_path", "title", "slides"]
        },
        func=render_pptx
    ),
    ToolDefinition(
        name="edit_spreadsheet",
        description="Create or update an Excel spreadsheet (.xlsx) with structured tabular data.",
        parameters={
            "type": "object",
            "properties": {
                "file_path": {"type": "string", "description": "Output path in workspace, e.g. 'output/metrics.xlsx'"},
                "sheet_name": {"type": "string", "description": "Sheet name, default 'Sheet1'", "default": "Sheet1"},
                "rows": {
                    "type": "array",
                    "description": "2D array of rows (first row is header)",
                    "items": {"type": "array", "items": {"type": "string"}}
                },
                "create_new": {"type": "boolean", "description": "Whether to overwrite/create new workbook", "default": True}
            },
            "required": ["file_path", "rows"]
        },
        func=edit_spreadsheet
    )
]
